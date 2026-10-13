"""Required numeric cells for the existing complete Excel monthly-report contract."""
from __future__ import annotations

import math
from contextlib import contextmanager
from pathlib import Path

import openpyxl

from .sku import SkuNormalizer


SHEET = "SKU销售汇总"
REQUIRED_FIELDS = {
    "units": 1,
    "productSales": 2,
    "settlementPayout": 3,
    "returns": 5,
    "landedCost": 8,
    "grossProfit": 9,
    "advertisingCost": 11,
    "storageCost": 12,
    "actualProfit": 13,
}


@contextmanager
def open_monthly_report(path: Path):
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as error:
        raise ValueError(f"月报无法读取: 文件={path}; 表={SHEET}; {error}") from error
    try:
        yield workbook
    finally:
        workbook.close()


def validated_monthly_rows(workbook, path: Path | str, normalizer: SkuNormalizer):
    """Yield individually validated SKU rows before aggregation.

    An uncached formula is None in data-only mode and is intentionally rejected.
    No business meaning is inferred from missing cells, including expense cells.
    """
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"月报缺少工作表: 文件={path}; 表={SHEET}")
    found = False
    for row_number, row in enumerate(workbook[SHEET].iter_rows(min_row=4, values_only=True), start=4):
        raw_sku = row[0] if row else None
        cleaned_sku = normalizer.clean(raw_sku)
        if not cleaned_sku or cleaned_sku.upper() in normalizer.ignore_values:
            continue
        match = normalizer.extract(raw_sku)
        if not match or len(match.bases) != 1:
            raise ValueError(
                f"月报 SKU 无法唯一映射: 文件={path}; 表={SHEET}; 行={row_number}; "
                f"SKU={raw_sku}; 字段=SKU"
            )
        found = True
        values = {}
        for field, index in REQUIRED_FIELDS.items():
            raw = row[index] if len(row) > index else None
            try:
                number = float(str(raw).replace(",", "").strip())
                if not math.isfinite(number):
                    raise ValueError("non-finite value")
            except (TypeError, ValueError, OverflowError):
                raise ValueError(
                    f"月报必要数值缺失或无效: 文件={path}; 表={SHEET}; 行={row_number}; "
                    f"SKU={raw_sku}; 字段={field}; 列={index + 1}; 值={raw!r}"
                ) from None
            values[field] = number
        yield match.bases[0], values
    if not found:
        raise ValueError(f"月报无有效数据: 文件={path}; 表={SHEET}")
