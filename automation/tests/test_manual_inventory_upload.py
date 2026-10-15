import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl

from store_ops.report_versions import current_reports, report_baseline_fingerprint
from store_ops.uploaded_data import inspect_batch, main, publish_batch, publish_target_fingerprint


class ManualInventoryUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.batch = self.root / "batch-manual123456"
        self.source = self.batch / "source" / "库存规划.xlsx"
        self.source.parent.mkdir(parents=True)
        self.reports = self.root / "reports"
        self.reports.mkdir()
        self.snapshots = self.root / "snapshots"
        self.write_workbook()
        row = {"sku": "MA001", "productName": "旧名称", "fbaSellable": 5, "awdAvailable": 4,
               "awdOutboundToFba": 2, "inTransitInventory": 7, "localInventory": 10, "pendingOrderQty": 2,
               "domesticSupplyTotal": 12, "dailySales": 1, "salesByMonth": [],
               "salesHistoryByMonth": [{"month": "2026-09", "units": 17}],
               "pendingOrders": [{"poNumber": "human-kept"}]}
        for market, filename in (("US", "inventory_dashboard.json"), ("CA", "inventory_dashboard.ca.json")):
            report = {"schemaVersion": 2, "market": market, "generatedAt": "2026-07-14T00:00:00Z",
                      "snapshots": {"fbaDate": "2026-07-14", "awdDate": "2026-06-26", "staleAfterDays": 14},
                      "summary": {}, "rows": [row, {**row, "sku": "MA003"}]}
            (self.reports / filename).write_text(json.dumps(report))
        self.manual = self.reports / "document_master.json"
        self.manual.write_text('{"shipmentHistory":[{"human":true}],"purchaseOrderLots":[{"human":true}]}')
        self.original_reports = {path.name: path.read_bytes() for path in self.reports.glob("*.json")}

    def tearDown(self):
        self.temp.cleanup()

    def write_workbook(self, *, replacements=None, missing=None, duplicate=False):
        headers = ["SKU", "品名", "工厂", "箱规", "工厂库存及已下订单", "FBA库存", "AWD备货", "FBA+在途库存", "最近月销售", "参考月销"]
        values = ["MA001", "实际名称", "工厂甲", 12, 30, 20, 8, 50, 60, 50]
        for name, value in (replacements or {}).items():
            values[headers.index(name)] = value
        if missing:
            index = headers.index(missing)
            del headers[index]
            del values[index]
        workbook = openpyxl.Workbook()
        us = workbook.active
        us.title = "库存规划"
        us.append(headers)
        us.append(values)
        if duplicate:
            us.append(values)
        if not replacements and not missing:
            us.append(["MA002", "新增产品", "工厂乙", 24, 10, 3, 2, 10, 30, 20])
        ca = workbook.create_sheet("加拿大库存计划 ")
        ca.append(["SKU", "品名", "工厂库存及已下订单", "FBA库存", "FBA+在途库存", "最近月销"])
        ca.append(["MA001", "加拿大", 999, 15, 45, 30])
        workbook.save(self.source)
        workbook.close()

    def inspect_actual(self):
        return inspect_batch(self.batch, confirmed_inventory_date="2026-10-12", inventory_scope="actual")

    def assert_original_reports(self):
        self.assertFalse((self.reports / "current.json").exists())
        self.assertEqual({path.name: path.read_bytes() for path in self.reports.glob("*.json")}, self.original_reports)

    def test_actual_inspect_then_guarded_publish_updates_confirmed_inventory_without_double_counting(self):
        original = self.source.read_bytes()
        preview = self.inspect_actual()
        self.assertEqual(preview["source"]["kind"], "manual-inventory-upload")
        self.assertEqual(preview["source"]["authentication"], "manual-upload")
        self.assertEqual(preview["source"]["inventoryScope"], "actual")
        self.assertEqual(preview["source"]["businessAsOf"], "2026-10-12")
        self.assertEqual(preview["source"]["confirmedInventoryDate"], "2026-10-12")
        self.assertEqual(preview["source"]["confirmation"], "explicit-user-confirmation")
        self.assertEqual(preview["source"]["sha256"], hashlib.sha256(original).hexdigest())
        self.assertNotIn("shareUrl", preview["source"])
        self.assertNotIn("fileToken", preview["source"])
        self.assert_original_reports()
        result = publish_batch(self.batch, self.reports, self.snapshots,
                               expected_baseline=report_baseline_fingerprint(self.reports),
                               expected_target=publish_target_fingerprint(self.batch))
        self.assertEqual(result["status"], "published")
        active = current_reports(self.reports)
        us = json.loads((active / "inventory_dashboard.json").read_text())
        rows = {row["sku"]: row for row in us["rows"]}
        self.assertEqual((rows["MA001"]["fbaSellable"], rows["MA001"]["awdAvailable"], rows["MA001"]["inTransitInventory"]), (20, 8, 22))
        self.assertEqual(rows["MA001"]["awdOutboundToFba"], 0)
        self.assertEqual((rows["MA001"]["domesticSupplyTotal"], rows["MA001"]["localInventory"], rows["MA001"]["pendingOrderQty"]), (30, 10, 2))
        self.assertFalse(rows["MA001"]["domesticSupplyBreakdownVerified"])
        self.assertEqual(rows["MA001"]["pendingOrders"], [{"poNumber": "human-kept"}])
        self.assertEqual(rows["MA001"]["salesHistoryByMonth"], [{"month": "2026-09", "units": 17}])
        self.assertEqual(rows["MA001"]["inventorySnapshotDate"], "2026-10-12")
        self.assertEqual(rows["MA002"]["inventorySnapshotDate"], "2026-10-12")
        self.assertEqual(rows["MA003"]["inventorySnapshotDate"], "2026-07-14")
        self.assertEqual(us["snapshots"]["fbaDate"], "2026-10-12")
        self.assertEqual(us["snapshots"]["awdDate"], "2026-10-12")
        self.assertTrue(us["snapshots"]["inTransitConfirmed"])
        self.assertTrue(us["snapshots"]["awdAvailableConfirmed"])
        self.assertFalse(us["snapshots"]["businessDateUnknown"])
        ca = json.loads((active / "inventory_dashboard.ca.json").read_text())
        self.assertEqual(ca["rows"][0]["inTransitInventory"], 30)
        self.assertEqual(ca["rows"][0]["domesticSupplyTotal"], 30)
        self.assertFalse(ca["snapshots"]["awdSourceAvailable"])
        self.assertEqual((active / "document_master.json").read_bytes(), self.original_reports["document_master.json"])
        self.assertEqual(self.source.read_bytes(), original)

    def test_cli_inspect_forwards_actual_scope_and_business_date(self):
        arguments = ["uploaded_data", "inspect", "--batch-dir", str(self.batch), "--inventory-scope", "actual", "--confirmed-inventory-date", "2026-10-12"]
        output = io.StringIO()
        with patch.object(sys, "argv", arguments), patch("store_ops.uploaded_data.load_config", return_value=None), redirect_stdout(output):
            self.assertEqual(main(), 0)
        self.assertEqual(json.loads(output.getvalue())["source"]["businessAsOf"], "2026-10-12")

    def test_cli_explicit_inventory_returns_safe_validation_json_without_traceback(self):
        config = Path(__file__).parents[1] / "config" / "project.json"
        for case, expected in (("missing_awd", "AWD备货"), ("invalid_date", "日期"), ("not_inventory", "库存工作簿")):
            with self.subTest(case=case):
                self.write_workbook(missing="AWD备货" if case == "missing_awd" else None)
                if case == "not_inventory":
                    workbook = openpyxl.Workbook()
                    workbook.active.title = "其他资料"
                    workbook.save(self.source)
                    workbook.close()
                    self.source.rename(self.source.with_name("其他资料.xlsx"))
                result = subprocess.run([sys.executable, "-m", "store_ops.uploaded_data", "inspect", "--config", str(config),
                                         "--batch-dir", str(self.batch), "--inventory-scope", "actual", "--confirmed-inventory-date",
                                         "2026-02-30" if case == "invalid_date" else "2026-10-12"],
                                        capture_output=True, text=True, check=False, env=os.environ.copy())
                self.assertNotEqual(result.returncode, 0)
                self.assertTrue(result.stdout.strip().startswith("{"), result.stderr)
                payload = json.loads(result.stdout)
                self.assertEqual(payload["code"], "MANUAL_INVENTORY_VALIDATION")
                self.assertIn(expected, payload["error"])
                self.assertNotIn("Traceback", result.stdout + result.stderr)
                self.assertNotIn(str(self.root), result.stdout)

    def test_cli_keeps_unexpected_configuration_errors_out_of_safe_validation_channel(self):
        arguments = ["uploaded_data", "inspect", "--batch-dir", str(self.batch), "--inventory-scope", "actual", "--confirmed-inventory-date", "2026-10-12"]
        output = io.StringIO()
        with patch.object(sys, "argv", arguments), patch("store_ops.uploaded_data.load_config", side_effect=ValueError("synthetic-private-config")), redirect_stdout(output):
            with self.assertRaisesRegex(ValueError, "synthetic-private-config"):
                main()
        self.assertEqual(output.getvalue(), "")

    def test_actual_ca_only_standard_sheet_is_recognized_without_inventory_filename(self):
        workbook = openpyxl.load_workbook(self.source)
        del workbook["库存规划"]
        source = self.source.with_name("download.xlsx")
        workbook.save(source)
        workbook.close()
        self.source.unlink()
        try:
            preview = self.inspect_actual()
        except ValueError as error:
            self.fail(f"标准 CA 工作表应识别为库存，不应被拒绝：{error}")
        self.assertEqual(preview["files"][0]["type"], "inventory")
        self.assertEqual(set(preview["files"][0]["preview"]["markets"]), {"CA"})
        publish_batch(self.batch, self.reports, self.snapshots)
        active = current_reports(self.reports)
        self.assertEqual((active / "inventory_dashboard.json").read_bytes(), self.original_reports["inventory_dashboard.json"])
        self.assertEqual(json.loads((active / "inventory_dashboard.ca.json").read_text())["rows"][0]["fbaSellable"], 15)

    def test_actual_ca_without_awd_does_not_add_old_awd_pools_to_confirmed_network(self):
        self.inspect_actual()
        publish_batch(self.batch, self.reports, self.snapshots)
        report = json.loads((current_reports(self.reports) / "inventory_dashboard.ca.json").read_text())
        row = report["rows"][0]
        self.assertEqual(row["fbaSellable"] + row["inTransitInventory"] + row["awdAvailable"] + row["awdOutboundToFba"], 45)
        self.assertEqual(row["eligibleInventoryPosition"], 45)
        self.assertFalse(report["snapshots"]["awdAvailableConfirmed"])
        self.assertFalse(report["snapshots"]["awdSourceAvailable"])
        self.assertEqual(row["awdOnHand"], 0)

    def test_actual_inspect_rejects_missing_invalid_or_mismatched_date(self):
        for arguments in ({"inventory_scope": "actual"}, {"inventory_scope": "actual", "confirmed_inventory_date": "2026-02-30"},
                          {"confirmed_inventory_date": "2026-10-12"}, {"inventory_scope": "planning", "confirmed_inventory_date": "2026-10-12"},
                          {"inventory_scope": "unexpected"}):
            with self.subTest(arguments=arguments), self.assertRaises(ValueError):
                inspect_batch(self.batch, **arguments)
        self.assertFalse((self.batch / "manifest.json").exists())

    def test_actual_inspect_rejects_non_inventory_or_multiple_workbooks(self):
        second = self.source.with_name("second.xlsx")
        shutil.copy2(self.source, second)
        with self.assertRaises(ValueError):
            self.inspect_actual()
        second.unlink()
        self.source.unlink()
        workbook = openpyxl.Workbook()
        workbook.active.title = "其他资料"
        workbook.save(second)
        workbook.close()
        with self.assertRaises(ValueError):
            self.inspect_actual()
        self.assertFalse((self.batch / "manifest.json").exists())

    def test_actual_inspect_rejects_incomplete_or_invalid_stock_before_ready_preview(self):
        cases = [{"missing": name} for name in ("SKU", "FBA库存", "FBA+在途库存", "AWD备货", "工厂库存及已下订单")]
        cases += [{"replacements": {name: value}} for name, value in (("FBA库存", None), ("AWD备货", -1),
                   ("FBA+在途库存", 27), ("FBA库存", 1.5), ("工厂库存及已下订单", "NaN"), ("AWD备货", "inf"), ("SKU", "invalid"))]
        cases.append({"duplicate": True})
        for arguments in cases:
            with self.subTest(arguments=arguments):
                self.write_workbook(**arguments)
                with self.assertRaises(ValueError):
                    self.inspect_actual()
                self.assertFalse((self.batch / "manifest.json").exists())

    def test_actual_publish_revalidates_manual_provenance_and_source_sha(self):
        preview = self.inspect_actual()
        manifest = self.batch / "manifest.json"
        changes = [("kind", "wps-api-download"), ("kind", "wps-browser-download"), ("authentication", "official-oauth"),
                   ("inventoryScope", "planning"), ("confirmedInventoryDate", "2026-02-30"),
                   ("businessAsOf", "2026-10-13"), ("confirmation", None), ("sha256", "0" * 64)]
        for key, value in changes:
            with self.subTest(field=key, value=value):
                changed = json.loads(json.dumps(preview))
                changed["source"][key] = value
                manifest.write_text(json.dumps(changed))
                with self.assertRaises(ValueError):
                    publish_batch(self.batch, self.reports, self.snapshots)
                self.assert_original_reports()

    def test_publish_revalidates_actual_rows_and_preserves_existing_confirmation_guard(self):
        preview = self.inspect_actual()
        expected_baseline = report_baseline_fingerprint(self.reports)
        expected_target = publish_target_fingerprint(self.batch)
        changed = json.loads(json.dumps(preview))
        changed["source"]["inventoryScope"] = "planning"
        changed["source"].pop("confirmedInventoryDate")
        changed["source"].pop("confirmation")
        changed["source"]["businessAsOf"] = None
        (self.batch / "manifest.json").write_text(json.dumps(changed))
        with self.assertRaisesRegex(ValueError, "变化"):
            publish_batch(self.batch, self.reports, self.snapshots, expected_baseline=expected_baseline, expected_target=expected_target)
        self.assert_original_reports()
        self.write_workbook(replacements={"FBA库存": None})
        changed = json.loads(json.dumps(preview))
        changed["files"][0]["sha256"] = changed["source"]["sha256"] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        (self.batch / "manifest.json").write_text(json.dumps(changed))
        with self.assertRaises(ValueError):
            publish_batch(self.batch, self.reports, self.snapshots)
        self.assert_original_reports()

    def test_explicit_planning_preserves_overseas_stock_and_uses_combined_domestic_supply(self):
        preview = inspect_batch(self.batch, inventory_scope="planning")
        self.assertEqual(preview["source"]["kind"], "manual-inventory-upload")
        self.assertEqual(preview["source"]["inventoryScope"], "planning")
        self.assertIsNone(preview["source"]["businessAsOf"])
        self.assertEqual(preview["source"]["sha256"], preview["files"][0]["sha256"])
        publish_batch(self.batch, self.reports, self.snapshots)
        current = json.loads((current_reports(self.reports) / "inventory_dashboard.json").read_text())
        row = current["rows"][0]
        self.assertEqual((row["fbaSellable"], row["awdAvailable"], row["inTransitInventory"]), (5, 4, 7))
        self.assertEqual((row["localInventory"], row["pendingOrderQty"], row["domesticSupplyTotal"]), (10, 2, 30))
        self.assertEqual(current["snapshots"]["fbaDate"], "2026-07-14")
        self.assertTrue(current["snapshots"]["businessDateUnknown"])
        self.assertEqual(len(current["rows"]), 2)

    def test_planning_allows_other_files_but_only_one_inventory_workbook(self):
        second = self.source.with_name("发货清单.xlsx")
        book = openpyxl.Workbook()
        book.active.title = "其他资料"
        book.save(second)
        book.close()
        preview = inspect_batch(self.batch, inventory_scope="planning")
        self.assertEqual(preview["source"]["sha256"], next(item["sha256"] for item in preview["files"] if item["type"] == "inventory"))
        shutil.copy2(self.source, self.source.with_name("库存规划副本.xlsx"))
        with self.assertRaises(ValueError):
            inspect_batch(self.batch, inventory_scope="planning")
        self.source.unlink()
        self.source.with_name("库存规划副本.xlsx").unlink()
        self.assertNotIn("source", inspect_batch(self.batch, inventory_scope="planning"))

    def test_legacy_inspect_keeps_its_existing_metadata_free_contract(self):
        self.assertNotIn("source", inspect_batch(self.batch))


if __name__ == "__main__":
    unittest.main()
