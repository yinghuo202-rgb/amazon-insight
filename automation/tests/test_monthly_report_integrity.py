import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from store_ops.config import ProjectConfig
from store_ops.db import StateDb
from store_ops.jobs.inventory_dashboard import _read_monthly_sales_reports
from store_ops.jobs.monthly_reports import run as import_reports
from store_ops.jobs.profitability_snapshot import _read_report, run as build_profitability
from store_ops.jobs.sales_history import run as refresh_sales_history
from store_ops.report_versions import current_reports, report_transaction
from store_ops.sku import SkuNormalizer
from store_ops.uploaded_data import inspect_batch, publish_batch


COMPLETE_ROW = ["MA001", 10, 200, 150, None, 2, None, None, 40, 110, None, -20, -5, 85]
FIELDS = {1: "units", 2: "productSales", 3: "settlementPayout", 5: "returns", 8: "landedCost",
          9: "grossProfit", 11: "advertisingCost", 12: "storageCost", 13: "actualProfit"}


def write_monthly(path, rows):
    workbook = openpyxl.Workbook()
    workbook.active.title = "SKU销售汇总"
    for _ in range(3):
        workbook.active.append([])
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


class MonthlyReportIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.path = self.root / "2026.7月-Measureman销售和毛利报告-US.xlsx"
        self.normalizer = SkuNormalizer(r"(MA\d{3})", frozenset())
        self.config = ProjectConfig(self.root / "config.json", self.root, self.root, self.root,
                                    r"(MA\d{3})", frozenset(), (), {"sales_monthly_folder": "."})

    def database(self):
        database = StateDb(self.root / "db" / "operations.sqlite3")
        database.init()
        self.addCleanup(database.close)
        return database

    def assert_invalid(self, callback, field, row=4):
        with self.assertRaises(ValueError) as failure:
            callback()
        for location in (self.path.name, "SKU销售汇总", str(row), "MA001", field):
            self.assertIn(location, str(failure.exception))

    def test_each_required_financial_field_rejects_missing_value(self):
        # Removing any required-cell validation must expose an invalid financial report.
        for index, field in FIELDS.items():
            with self.subTest(field=field):
                row = COMPLETE_ROW.copy()
                row[index] = None
                write_monthly(self.path, [row])
                self.assert_invalid(lambda: _read_report(self.path, "US", "2026-07", self.normalizer), field)

    def test_invalid_numeric_values_and_uncached_formula_are_not_zero(self):
        for value in ("", "   ", "-", "not-numeric", "NaN", "Infinity", "-Infinity", "=1+1"):
            with self.subTest(value=value):
                row = COMPLETE_ROW.copy()
                row[11] = value
                write_monthly(self.path, [row])
                self.assert_invalid(lambda: _read_report(self.path, "US", "2026-07", self.normalizer), "advertisingCost")

    def test_explicit_zero_and_existing_financial_formulas_are_preserved(self):
        zero = ["MA002", 0, "0", 0, None, 0, None, None, 0, 0, None, 0, 0, 0]
        write_monthly(self.path, [COMPLETE_ROW, zero])
        actual, empty = _read_report(self.path, "US", "2026-07", self.normalizer)
        self.assertEqual((actual["netUnits"], actual["averagePrice"], actual["advertisingCost"],
                          actual["storageCost"], actual["conservativeMargin"]), (8, 20, 20, 5, 0.4))
        self.assertEqual([empty[field] for field in FIELDS.values()], [0] * len(FIELDS))
        self.assertIsNone(empty["averagePrice"])

    def test_duplicate_sku_missing_row_is_not_partially_aggregated(self):
        missing = COMPLETE_ROW.copy()
        missing[5] = None
        write_monthly(self.path, [COMPLETE_ROW, missing])
        self.assert_invalid(lambda: _read_report(self.path, "US", "2026-07", self.normalizer), "returns", row=5)

    def test_configured_total_row_with_blank_expenses_is_not_a_sku(self):
        total = ["总计", 10, 200]
        write_monthly(self.path, [COMPLETE_ROW, total])
        normalizer = SkuNormalizer(r"(MA\d{3})", frozenset({"总计"}))
        try:
            rows = _read_report(self.path, "US", "2026-07", normalizer)
        except ValueError as error:
            self.fail(f"Configured ignored total row is not SKU data: {error}")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["productSales"], 200)

    def test_unmapped_or_ambiguous_sku_is_not_silently_discarded(self):
        for raw_sku in ("unknown-product", "MA001 / MA002"):
            with self.subTest(sku=raw_sku):
                row = COMPLETE_ROW.copy()
                row[0] = raw_sku
                write_monthly(self.path, [row])
                with self.assertRaises(ValueError) as failure:
                    _read_report(self.path, "US", "2026-07", self.normalizer)
                for location in (self.path.name, "SKU销售汇总", "行=4", raw_sku, "SKU"):
                    self.assertIn(location, str(failure.exception))

    def test_preview_uses_configured_ignored_total_and_rejects_bad_sku_mapping(self):
        batch = self.root / "batch-sku-map"
        source = batch / "source"
        source.mkdir(parents=True)
        self.path = source / self.path.name
        write_monthly(self.path, [COMPLETE_ROW, ["总计", 10, 200]])
        self.assertTrue(inspect_batch(batch)["files"][0]["publishable"])
        for raw_sku in ("unknown-product", "MA001 / MA002"):
            with self.subTest(sku=raw_sku):
                row = COMPLETE_ROW.copy()
                row[0] = raw_sku
                write_monthly(self.path, [row])
                preview = inspect_batch(batch)["files"][0]
                self.assertFalse(preview["publishable"])
                self.assertIn(raw_sku, preview["error"])

    def test_profitability_failure_preserves_previous_report(self):
        reports = self.root / "reports"
        reports.mkdir()
        prior = b'{"rows": [{"market": "US", "reportMonth": "2026-06", "sku": "MA001", "actualProfit": 999}], "sources": []}'
        (reports / "profitability.json").write_bytes(prior)
        row = COMPLETE_ROW.copy()
        row[8] = None
        write_monthly(self.path, [row])
        with patch.dict(os.environ, {"STORE_OPS_PROFITABILITY_ROOT": str(self.root)}):
            self.assert_invalid(lambda: build_profitability(self.config, self.database()), "landedCost")
        self.assertEqual((reports / "profitability.json").read_bytes(), prior)

    def test_import_validates_all_files_before_copying_any(self):
        write_monthly(self.path, [COMPLETE_ROW])
        second = self.root / "2026.8月-Measureman销售和毛利报告-US.xlsx"
        row = COMPLETE_ROW.copy()
        row[2] = None
        write_monthly(second, [row])
        with self.assertRaises(ValueError):
            import_reports(self.config, self.database(), [self.path, second])
        self.assertFalse((self.root / "incoming").exists())

    def test_invalid_monthly_preview_blocks_mixed_batch_and_keeps_pointer(self):
        batch = self.root / "batch-mixed"
        source = batch / "source"
        source.mkdir(parents=True)
        self.path = source / self.path.name
        row = COMPLETE_ROW.copy()
        row[12] = None
        write_monthly(self.path, [row])
        inventory = openpyxl.Workbook()
        inventory.active.title = "库存规划"
        inventory.active.append(["SKU", "FBA库存", "工厂库存及已下订单", "参考月销"])
        inventory.active.append(["MA001", 10, 20, 30])
        inventory.save(source / "库存规划.xlsx")
        inventory.close()
        reports = self.root / "reports"
        snapshots = self.root / "snapshots"
        with report_transaction(reports, snapshots) as (stage, _):
            (stage / "inventory_dashboard.json").write_text('{"rows": [{"sku": "MA001", "fbaSellable": 5}]}')
        pointer = (reports / "current.json").read_bytes()
        prior = (current_reports(reports) / "inventory_dashboard.json").read_bytes()
        preview = inspect_batch(batch)
        monthly = next(item for item in preview["files"] if item["name"] == self.path.name)
        self.assertFalse(monthly["publishable"])
        self.assertIn("storageCost", monthly["error"])
        self.assert_invalid(lambda: publish_batch(batch, reports, snapshots), "storageCost")
        self.assertEqual((reports / "current.json").read_bytes(), pointer)
        self.assertEqual((current_reports(reports) / "inventory_dashboard.json").read_bytes(), prior)

    def test_publish_revalidates_monthly_even_if_manifest_marks_it_unpublishable(self):
        batch = self.root / "batch-stale"
        source = batch / "source"
        source.mkdir(parents=True)
        self.path = source / self.path.name
        row = COMPLETE_ROW.copy()
        row[13] = None
        write_monthly(self.path, [row])
        inspect_batch(batch)
        manifest = json.loads((batch / "manifest.json").read_text())
        manifest["files"][0].pop("error", None)
        manifest["files"][0]["publishable"] = False
        (batch / "manifest.json").write_text(json.dumps(manifest))
        reports = self.root / "reports"
        reports.mkdir()
        (reports / "kept.json").write_text('{"old": true}')
        self.assert_invalid(lambda: publish_batch(batch, reports, self.root / "snapshots"), "actualProfit")
        self.assertFalse((reports / "current.json").exists())

    def test_history_checks_latest_source_not_superseded_invalid_duplicate(self):
        old = self.root / "old"
        old.mkdir()
        row = COMPLETE_ROW.copy()
        row[2] = None
        old_path = old / self.path.name
        write_monthly(old_path, [row])
        write_monthly(self.path, [COMPLETE_ROW])
        os.utime(old_path, (100, 100))
        os.utime(self.path, (200, 200))
        _, months, actual, paths, _ = _read_monthly_sales_reports(self.root, "US", 3, self.normalizer)
        self.assertEqual(months, ["2026-07"])
        self.assertEqual(actual["2026-07"]["units"], 10)
        self.assertEqual(paths, [self.path])

    def test_selected_latest_corrupt_source_is_not_silently_skipped(self):
        old = self.root / "old"
        old.mkdir()
        old_path = old / self.path.name
        write_monthly(old_path, [COMPLETE_ROW])
        self.path.write_bytes(b"not an Excel workbook")
        os.utime(old_path, (100, 100))
        os.utime(self.path, (200, 200))
        with self.assertRaises(ValueError) as failure:
            _read_monthly_sales_reports(self.root, "US", 3, self.normalizer)
        self.assertIn(str(self.path), str(failure.exception))

    def test_newer_unrelated_inventory_file_cannot_replace_monthly_source(self):
        write_monthly(self.path, [COMPLETE_ROW])
        unrelated = self.root / "2026.7月-库存规划.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "库存规划"
        workbook.active.append(["SKU", "FBA库存"])
        workbook.active.append(["MA001", 999])
        workbook.save(unrelated)
        workbook.close()
        os.utime(self.path, (100, 100))
        os.utime(unrelated, (200, 200))
        try:
            _, months, actual, paths, _ = _read_monthly_sales_reports(self.root, "US", 3, self.normalizer)
        except ValueError as error:
            self.fail(f"Unrelated inventory workbook must not be selected as a monthly source: {error}")
        self.assertEqual(months, ["2026-07"])
        self.assertEqual(actual["2026-07"]["units"], 10)
        self.assertEqual(paths, [self.path])

    def test_independent_sales_history_failure_keeps_previous_report(self):
        reports = self.root / "reports"
        reports.mkdir()
        prior = b'{"rows": [{"sku": "MA001"}], "sales": {"historyMonths": []}}'
        (reports / "inventory_dashboard.json").write_bytes(prior)
        row = COMPLETE_ROW.copy()
        row[1] = None
        write_monthly(self.path, [row])
        self.assert_invalid(lambda: refresh_sales_history(self.config, self.database()), "units")
        self.assertEqual((reports / "inventory_dashboard.json").read_bytes(), prior)

    def test_later_market_validation_failure_keeps_all_previous_history_reports(self):
        from dataclasses import replace
        reports = self.root / "reports"
        reports.mkdir()
        prior = b'{"rows": [{"sku": "MA001"}], "sales": {"historyMonths": []}}'
        for filename in ("inventory_dashboard.json", "inventory_dashboard.ca.json"):
            (reports / filename).write_bytes(prior)
        write_monthly(self.path, [COMPLETE_ROW])
        canada = self.root / "2026.7月-Measureman销售和毛利报告-CA.xlsx"
        row = COMPLETE_ROW.copy()
        row[5] = None
        write_monthly(canada, [row])
        config = replace(self.config, inventory_dashboard={"sales_monthly_folder": ".", "markets": {"CA": {}}})
        with self.assertRaises(ValueError):
            refresh_sales_history(config, self.database())
        for filename in ("inventory_dashboard.json", "inventory_dashboard.ca.json"):
            self.assertEqual((reports / filename).read_bytes(), prior)
